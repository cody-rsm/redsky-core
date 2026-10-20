import io
import random
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

# Every document here is synthetic. Names, figures and entities are invented
# so the dataset can live in the repository and in CI logs.
COMPANY = "Northwind Test Holdings"

LETTER_PAGES = [
    [
        f"{COMPANY}",
        "Quarterly Investor Letter",
        "",
        "Dear partners,",
        "Revenue for the quarter was 4.2 million dollars,",
        "up twelve percent from the prior year.",
        "Gross margin improved to 61 percent.",
        "We opened a second facility in Ohio.",
    ],
    [
        "Outlook",
        "",
        "We expect the Series B to close in November.",
        "Headcount will grow from 48 to 60 people.",
        "The board meets again on the fourth of December.",
        "Sincerely, the finance team",
    ],
]

TABLE_ROWS = [
    ["Metric", "2025", "2026"],
    ["Revenue", "3,750", "4,200"],
    ["Gross margin", "58%", "61%"],
    ["Headcount", "41", "48"],
]

MEMO_PARAGRAPHS = [
    "This memorandum summarises the proposed investment in a synthetic test company.",
    "The company sells industrial sensors to utilities and reported steady growth.",
    "Key risks include customer concentration and supply chain delays.",
]


@dataclass
class SampleDocument:
    name: str
    path: Path
    expected_words: list[str] = field(default_factory=list)
    expected_pages: int | None = None
    expected_tables: int = 0


def _font(size: int) -> ImageFont.ImageFont:
    for candidate in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def render_scan(lines: list[str], seed: int, dpi: int = 300) -> Image.Image:
    # US Letter at scanner resolution, with a slight tilt, blur and speckle so
    # OCR has something like a real scan to work on.
    width, height = int(8.5 * dpi), int(11 * dpi)
    image = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(image)
    font = _font(int(dpi * 0.16))
    y = int(dpi * 1.0)
    for line in lines:
        if line:
            draw.text((int(dpi * 1.0), y), line, fill=20, font=font)
        y += int(dpi * 0.3)
    generator = random.Random(seed)
    pixels = image.load()
    for _ in range(width * height // 400):
        pixels[generator.randrange(width), generator.randrange(height)] = generator.choice((0, 200))
    image = image.rotate(0.6, fillcolor=255, resample=Image.Resampling.BICUBIC)
    return image.filter(ImageFilter.GaussianBlur(0.6))


def _words(lines: list[str]) -> list[str]:
    return [word.strip(".,") for line in lines for word in line.split() if len(word) > 3]


def make_native_pdf(path: Path) -> SampleDocument:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    table = Table(TABLE_ROWS)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    story = [
        Paragraph(f"{COMPANY} Investment Memorandum", styles["Title"]),
        Paragraph("Summary", styles["Heading1"]),
        *[Paragraph(text, styles["BodyText"]) for text in MEMO_PARAGRAPHS],
        Spacer(1, 12),
        Paragraph("Financial summary", styles["Heading2"]),
        table,
        PageBreak(),
        Paragraph("Recommendation", styles["Heading1"]),
        Paragraph("The deal team recommends proceeding to confirmatory diligence.", styles["BodyText"]),
    ]
    SimpleDocTemplate(str(path), pagesize=letter, title="Synthetic memo").build(story)
    return SampleDocument(
        "native_pdf",
        path,
        expected_words=["Memorandum", "industrial", "sensors", "Revenue", "Recommendation", "diligence"],
        expected_pages=2,
        expected_tables=1,
    )


def make_scanned_pdf(path: Path) -> SampleDocument:
    images = [render_scan(lines, seed=index) for index, lines in enumerate(LETTER_PAGES)]
    images[0].save(path, format="PDF", resolution=300, save_all=True, append_images=images[1:])
    return SampleDocument(
        "scanned_pdf",
        path,
        expected_words=_words(LETTER_PAGES[0] + LETTER_PAGES[1]),
        expected_pages=2,
    )


def make_mixed_pdf(path: Path) -> SampleDocument:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setFont("Helvetica", 12)
    y = 700
    for text in MEMO_PARAGRAPHS:
        pdf.drawString(72, y, text)
        y -= 20
    pdf.showPage()
    scan = render_scan(LETTER_PAGES[1], seed=7, dpi=200)
    buffer = io.BytesIO()
    scan.save(buffer, format="PNG")
    buffer.seek(0)
    pdf.drawImage(ImageReader(buffer), 0, 0, width=letter[0], height=letter[1])
    pdf.showPage()
    pdf.save()
    return SampleDocument(
        "mixed_pdf",
        path,
        expected_words=["memorandum", "sensors", "Outlook", "Series", "November", "Headcount"],
        expected_pages=2,
    )


def make_png(path: Path) -> SampleDocument:
    render_scan(LETTER_PAGES[0], seed=3).save(path, format="PNG")
    return SampleDocument("png", path, expected_words=_words(LETTER_PAGES[0]), expected_pages=1)


def make_jpeg(path: Path) -> SampleDocument:
    render_scan(LETTER_PAGES[1], seed=4).convert("RGB").save(path, format="JPEG", quality=85)
    return SampleDocument("jpeg", path, expected_words=_words(LETTER_PAGES[1]), expected_pages=1)


def make_tiff(path: Path) -> SampleDocument:
    frames = [render_scan(lines, seed=10 + index, dpi=200) for index, lines in enumerate(LETTER_PAGES)]
    frames[0].save(
        path, format="TIFF", save_all=True, append_images=frames[1:], dpi=(200, 200), compression="tiff_deflate"
    )
    return SampleDocument(
        "tiff", path, expected_words=_words(LETTER_PAGES[0] + LETTER_PAGES[1]), expected_pages=2
    )


def make_docx(path: Path) -> SampleDocument:
    from docx import Document

    document = Document()
    document.add_heading(f"{COMPANY} Board Update", 1)
    for text in MEMO_PARAGRAPHS:
        document.add_paragraph(text)
    document.add_heading("Financial summary", 2)
    table = document.add_table(rows=len(TABLE_ROWS), cols=len(TABLE_ROWS[0]))
    for row_index, row in enumerate(TABLE_ROWS):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(path)
    return SampleDocument(
        "docx",
        path,
        expected_words=["Board", "Update", "sensors", "Revenue", "Headcount"],
        expected_tables=1,
    )


def make_pptx(path: Path) -> SampleDocument:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = f"{COMPANY}"
    title_slide.placeholders[1].text = "Series B pitch deck"
    bullets = presentation.slides.add_slide(presentation.slide_layouts[1])
    bullets.shapes.title.text = "Why now"
    bullets.placeholders[1].text = "Utilities are replacing legacy sensors\nRegulation requires monitoring"
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Traction"
    shape = table_slide.shapes.add_table(
        len(TABLE_ROWS), len(TABLE_ROWS[0]), Inches(1), Inches(2), Inches(6), Inches(2)
    )
    for row_index, row in enumerate(TABLE_ROWS):
        for column_index, value in enumerate(row):
            shape.table.cell(row_index, column_index).text = value
    presentation.save(path)
    return SampleDocument(
        "pptx",
        path,
        expected_words=["pitch", "Utilities", "Regulation", "Traction", "Revenue"],
        expected_pages=3,
        expected_tables=1,
    )


def make_xlsx(path: Path) -> SampleDocument:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    for row in TABLE_ROWS:
        sheet.append(row)
    cash = workbook.create_sheet("Cash")
    cash.append(["Month", "Cash balance"])
    for month, balance in (("July", 1200), ("August", 1350), ("September", 1410)):
        cash.append([month, balance])
    workbook.save(path)
    return SampleDocument(
        "xlsx",
        path,
        expected_words=["Revenue", "Headcount", "September"],
        expected_tables=2,
    )


def make_blank_pdf(path: Path) -> SampleDocument:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.showPage()
    pdf.showPage()
    pdf.save()
    return SampleDocument("blank_pdf", path, expected_pages=2)


def make_encrypted_pdf(path: Path) -> SampleDocument:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=letter, encrypt="synthetic-password")
    pdf.drawString(72, 700, "Locked synthetic document")
    pdf.showPage()
    pdf.save()
    return SampleDocument("encrypted_pdf", path)


def make_corrupt_pdf(path: Path) -> SampleDocument:
    path.write_bytes(b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R >>\nthis is not a pdf body")
    return SampleDocument("corrupt_pdf", path)


def make_macro_docx(path: Path) -> SampleDocument:
    make_docx(path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml",
                    b"ms-word.document.macroEnabled.main+xml",
                ).replace(
                    b"application/vnd.openxmlformats-officedocument.ms-word",
                    b"application/vnd.ms-word",
                )
            target.writestr(entry, data)
        target.writestr("word/vbaProject.bin", b"\x00" * 64)
    path.write_bytes(buffer.getvalue())
    return SampleDocument("macro_docx", path)


def make_zip_bomb(path: Path) -> SampleDocument:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"\x00" * (20 * 1024 * 1024))
    return SampleDocument("zip_bomb", path)


def make_pixel_bomb(path: Path) -> SampleDocument:
    Image.new("1", (12_000, 12_000), 1).save(path, format="PNG")
    return SampleDocument("pixel_bomb", path)


GOOD_DOCUMENTS = {
    "native.pdf": make_native_pdf,
    "scanned.pdf": make_scanned_pdf,
    "mixed.pdf": make_mixed_pdf,
    "letter.png": make_png,
    "letter.jpg": make_jpeg,
    "fax.tiff": make_tiff,
    "board-update.docx": make_docx,
    "pitch-deck.pptx": make_pptx,
    "financials.xlsx": make_xlsx,
}

BAD_DOCUMENTS = {
    "blank.pdf": make_blank_pdf,
    "encrypted.pdf": make_encrypted_pdf,
    "corrupt.pdf": make_corrupt_pdf,
    "macro.docx": make_macro_docx,
    "bomb.docx": make_zip_bomb,
    "pixel-bomb.png": make_pixel_bomb,
}


def generate(directory: Path, include_bad: bool = True) -> dict[str, SampleDocument]:
    directory.mkdir(parents=True, exist_ok=True)
    builders = {**GOOD_DOCUMENTS, **(BAD_DOCUMENTS if include_bad else {})}
    return {name: builder(directory / name) for name, builder in builders.items()}


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1] if len(sys.argv) > 1 else "sample-documents")
    for sample in generate(target).values():
        print(f"{sample.name:15} {sample.path} ({sample.path.stat().st_size} bytes)")
